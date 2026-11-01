"""Investigation report (.docx) builder — deterministic rendering, no LLM calls."""

import io
import json
import uuid

from docx import Document
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.db.models.investigation import Investigation


def _result(record: Investigation) -> dict:
    try:
        parsed = json.loads(record.result_json or "{}")
    except ValueError:
        parsed = {}
    return parsed if isinstance(parsed, dict) else {}


def _evidence(record: Investigation) -> list:
    try:
        items = json.loads(record.evidence_snapshot or "[]")
    except ValueError:
        items = []
    return items if isinstance(items, list) else []


def _runbooks(record: Investigation) -> list:
    try:
        items = json.loads(getattr(record, "runbook_snapshot", None) or "[]")
    except ValueError:
        items = []
    return items if isinstance(items, list) else []


def build_report_docx(session: Session, investigation_id: uuid.UUID) -> tuple[bytes, str]:
    """Render the investigation report. Returns (docx bytes, filename)."""
    record = session.get(Investigation, investigation_id)
    if record is None:
        raise KeyError(f"investigation not found: {investigation_id}")
    result = _result(record)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading("DeployIQ Investigation Report", level=1)
    doc.add_paragraph(
        f"Investigation {str(record.id)[:8]} · Incident {str(record.incident_id)[:8]} · "
        f"Status {record.status}"
    )
    meta = doc.add_paragraph()
    meta.add_run(
        f"Model: {record.actual_model or record.requested_model or '—'} · "
        f"Tokens: {record.total_tokens} · Duration: {(record.duration_ms or 0) / 1000:.1f}s · "
        f"Prompt {record.prompt_version}"
    )

    doc.add_heading("Problem summary", level=2)
    doc.add_paragraph(str(result.get("explanation") or "No explanation recorded."))

    doc.add_heading("Root cause", level=2)
    doc.add_paragraph(str(result.get("root_cause") or "Not established."))

    doc.add_heading("Evidence used", level=2)
    evidence = _evidence(record)
    if not evidence:
        doc.add_paragraph("No evidence snapshot recorded.")
    for entry in evidence:
        if isinstance(entry, dict):
            doc.add_paragraph(
                f"{entry.get('id', '?')} ({entry.get('kind', '?')}): {entry.get('summary', '')}",
                style="List Bullet",
            )

    doc.add_heading("Findings", level=2)
    findings = result.get("findings") or []
    if not findings:
        doc.add_paragraph("No findings recorded.")
    for finding in findings:
        if not isinstance(finding, dict):
            continue
        doc.add_paragraph(str(finding.get("description", "")), style="List Bullet")
        evidence_ids = ", ".join(str(e) for e in finding.get("evidence_ids", []) or [])
        runbook_ids = ", ".join(str(e) for e in finding.get("runbook_ids", []) or [])
        if evidence_ids or runbook_ids:
            doc.add_paragraph(f"Cites evidence [{evidence_ids}] runbooks [{runbook_ids}]")

    doc.add_heading("Recommendations (require approval)", level=2)
    recommendations = result.get("recommendations") or []
    if not recommendations:
        doc.add_paragraph("No recommendations recorded.")
    for rec in recommendations:
        if not isinstance(rec, dict):
            continue
        doc.add_paragraph(
            f"{rec.get('title', '')} [risk: {rec.get('risk_level', '?')}, "
            f"approval required: {bool(rec.get('requires_approval', True))}]",
            style="List Bullet",
        )
        for step in rec.get("steps", []) or []:
            doc.add_paragraph(str(step), style="List Bullet 2")

    doc.add_heading("Verification steps", level=2)
    verification = result.get("verification_steps") or []
    if not verification:
        doc.add_paragraph("No verification steps recorded.")
    for step in verification:
        doc.add_paragraph(str(step), style="List Number")

    limitations = str(result.get("limitations") or "").strip()
    if limitations:
        doc.add_heading("Limitations", level=2)
        doc.add_paragraph(limitations)

    buffer = io.BytesIO()
    doc.save(buffer)
    filename = f"deployiq-{str(record.incident_id)[:8]}-{str(record.id)[:8]}-report.docx"
    return buffer.getvalue(), filename
